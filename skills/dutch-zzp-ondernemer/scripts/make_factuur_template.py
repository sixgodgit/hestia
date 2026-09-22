#!/usr/bin/env python3
"""Generate a Dutch invoice (.docx) holding every mandatory btw field.

Run with the agent venv python -- python-docx is installed there, not in the
system python3:

    python make_factuur_template.py [OUT.docx]

Edit CONFIG to the user's registered details first (btw-id, handelsnaam, KVK),
then run it; the file lands with `[placeholders]` for anything genuinely unknown
(KVK-nummer, IBAN, client, amounts). Mandatory field list and the reasoning live
in references/zzp-checklist-nl.md section 3.

Never print the omzetbelastingnummer here -- it embeds the BSN and must stay off
invoices and websites.
"""
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

CONFIG = {
    "handelsnaam": "[Handelsnaam]",
    "naam": "[Volledige naam]",
    "adres": "[Straat + huisnummer]",
    "postcode_plaats": "[Postcode + plaats]",
    "btw_id": "[NL123456789B01]",
    "kvk": "[KVK-nummer]",
    "iban": "[IBAN]",
    "contact": "[telefoon]  ·  [e-mail]",
    "betalingstermijn": "14 dagen na factuurdatum",
    "btw_tarieven": ["btw 21%", "btw 9%"],
}


def build(out, c):
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Cm(2)
    sec.left_margin = sec.right_margin = Cm(2.2)
    doc.styles["Normal"].font.size = Pt(10.5)

    def para(text="", bold=False, size=10.5, align=None, space_after=2, italic=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(space_after)
        if align is not None:
            p.alignment = align
        r = p.add_run(text)
        r.bold, r.italic = bold, italic
        r.font.size = Pt(size)
        return p

    right = WD_ALIGN_PARAGRAPH.RIGHT
    para(c["handelsnaam"], bold=True, size=14, align=right, space_after=0)
    for line in (c["naam"], c["adres"], c["postcode_plaats"], c["contact"]):
        para(line, align=right, space_after=0)
    para(f"KVK: {c['kvk']}  ·  btw-id: {c['btw_id']}", align=right, space_after=0)
    para(f"IBAN: {c['iban']}", align=right, space_after=14)

    para("FACTUUR", bold=True, size=16, space_after=12)
    para("Aan:", bold=True, space_after=0)
    for line in ("[Naam klant]", "[Adres klant]", "[Postcode + plaats klant]"):
        para(line, space_after=0)
    para("[btw-id klant, indien van toepassing]", italic=True, space_after=14)

    meta = doc.add_table(rows=4, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, (k, v) in enumerate([
        ("Factuurnummer:", "[<jaar>-0001]"),
        ("Factuurdatum:", "[dd-mm-jjjj]"),
        ("Leverdatum / periode:", "[dd-mm-jjjj]"),
        ("Betalingstermijn:", c["betalingstermijn"]),
    ]):
        meta.cell(i, 0).text, meta.cell(i, 1).text = k, v
        for cell in meta.rows[i].cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                for r in p.runs:
                    r.font.size = Pt(10.5)
        meta.cell(i, 0).paragraphs[0].runs[0].bold = True
    para("", space_after=10)

    tbl = doc.add_table(rows=4, cols=4)
    tbl.style = "Table Grid"
    for j, h in enumerate(["Omschrijving", "Aantal", "Tarief excl. btw", "Bedrag excl. btw"]):
        cell = tbl.cell(0, j)
        cell.text = h
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(10.5)
    tbl.cell(1, 0).text = "[omschrijving werkzaamheden / product]"
    tbl.cell(1, 1).text = "[1]"
    tbl.cell(1, 2).text = tbl.cell(1, 3).text = "[€ 0,00]"
    tbl.cell(2, 0).text = "[eventueel regel 2]"
    tbl.cell(2, 2).text = tbl.cell(2, 3).text = "[€ 0,00]"
    tbl.cell(3, 0).text = "Subtotaal excl. btw"
    tbl.cell(3, 3).text = "[€ 0,00]"
    tbl.cell(3, 0).paragraphs[0].runs[0].bold = True
    para("", space_after=8)

    tot = doc.add_table(rows=len(c["btw_tarieven"]) + 1, cols=2)
    tot.alignment = WD_TABLE_ALIGNMENT.RIGHT
    rows = [(t, "[€ 0,00]") for t in c["btw_tarieven"]] + [("Totaal te betalen", "[€ 0,00]")]
    for i, (k, v) in enumerate(rows):
        tot.cell(i, 0).text, tot.cell(i, 1).text = k, v
        last = i == len(rows) - 1
        for cell in tot.rows[i].cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                for r in p.runs:
                    r.bold = last
                    r.font.size = Pt(11.5 if last else 10.5)
    para("", space_after=12)

    para(
        f"Gelieve het totaalbedrag binnen {c['betalingstermijn'].split()[0]} dagen te voldoen op "
        f"IBAN {c['iban']} t.n.v. {c['handelsnaam']}, onder vermelding van het factuurnummer.",
        size=9.5,
        space_after=2,
    )
    para(
        f"{c['handelsnaam']} · {c['naam']} · {c['adres']}, {c['postcode_plaats']} "
        f"· KVK {c['kvk']} · btw-id {c['btw_id']}",
        size=9,
        italic=True,
        space_after=2,
    )
    para(
        "Intracommunautaire levering (0% btw): VIES-controle op factuurdatum bewaren.",
        size=8.5,
        italic=True,
        space_after=0,
    )
    doc.save(out)
    return out


if __name__ == "__main__":
    target = sys.argv[1] if len(sys.argv) > 1 else "factuur_template.docx"
    print("geschreven:", build(target, CONFIG))
